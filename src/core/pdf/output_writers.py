"""
Write translated PDF text to txt, markdown, or docx output files.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable, List, Union

from docx import Document

_ALL_CAPS = re.compile(r"^[A-Z0-9][A-Z0-9\s\-:,.']{2,}$")
_SUPPORTED_OUTPUT_EXTENSIONS = {".txt", ".md", ".docx"}


def _normalize_paragraphs(paragraphs: Iterable[str]) -> List[str]:
    return [p.strip() for p in paragraphs if (p or "").strip()]


def _looks_like_heading(paragraph: str) -> bool:
    text = paragraph.strip()
    if not text or len(text) > 120:
        return False
    if _ALL_CAPS.match(text):
        return True
    if text.endswith(":") and len(text.split()) <= 12:
        return True
    return False


def write_txt(paragraphs: Iterable[str], output_path: Union[str, Path]) -> None:
    """Write plain UTF-8 text with blank lines between paragraphs."""
    path = Path(output_path)
    body = "\n\n".join(_normalize_paragraphs(paragraphs))
    path.write_text(body, encoding="utf-8")


def write_markdown(paragraphs: Iterable[str], output_path: Union[str, Path]) -> None:
    """Write markdown with lightweight heading heuristics."""
    path = Path(output_path)
    lines: List[str] = []
    for para in _normalize_paragraphs(paragraphs):
        if _looks_like_heading(para):
            lines.append(f"## {para}")
        else:
            lines.append(para)
        lines.append("")
    path.write_text("\n".join(lines).rstrip() + "\n", encoding="utf-8")


def write_docx(paragraphs: Iterable[str], output_path: Union[str, Path]) -> None:
    """Write a simple DOCX with one paragraph per block."""
    path = Path(output_path)
    document = Document()
    normalized = _normalize_paragraphs(paragraphs)
    if not normalized:
        document.add_paragraph("")
    else:
        for para in normalized:
            if _looks_like_heading(para):
                document.add_heading(para, level=2)
            else:
                document.add_paragraph(para)
    document.save(str(path))


def resolve_output_extension(output_path: Union[str, Path]) -> str:
    """Return a supported output extension for a PDF translation job."""
    ext = Path(output_path).suffix.lower()
    if ext in _SUPPORTED_OUTPUT_EXTENSIONS:
        return ext
    return ".txt"


def write_translated_output(
    paragraphs: Iterable[str],
    output_path: Union[str, Path],
    *,
    output_extension: str | None = None,
) -> str:
    """
    Write translated paragraphs to the requested output format.

    Returns:
        The extension actually used (``.txt``, ``.md``, or ``.docx``).
    """
    path = Path(output_path)
    ext = (output_extension or path.suffix or ".txt").lower()
    if not ext.startswith("."):
        ext = f".{ext}"
    if ext not in _SUPPORTED_OUTPUT_EXTENSIONS:
        ext = ".txt"
        path = path.with_suffix(ext)

    if ext == ".md":
        write_markdown(paragraphs, path)
    elif ext == ".docx":
        write_docx(paragraphs, path)
    else:
        write_txt(paragraphs, path)
    return ext
